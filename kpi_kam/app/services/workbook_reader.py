"""Shared read-only workbook access preserving original Decimal XML tokens."""
from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from posixpath import join, normpath
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from app.services.excel_importer import _numeric_cells, _cell_value, ExcelImportError, N, REL


@dataclass
class WorkbookData:
    source_file: str
    source_hash: str
    sheets: dict
    formulas_without_cache: dict
    merged_spans: dict


def read_workbook(path):
    path=Path(path)
    if path.suffix.casefold() not in ('.xlsx','.xlsm'):
        raise ExcelImportError('Поддерживаются .xlsx и .xlsm')
    workbook=None
    try:
        payload=path.read_bytes()
        workbook=load_workbook(BytesIO(payload),read_only=True,data_only=True,
                               keep_links=False,keep_vba=False)
        sheets,uncached,merged={},{},{}
        with ZipFile(BytesIO(payload)) as archive:
            document=ET.fromstring(archive.read('xl/workbook.xml'))
            links=ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
            for sheet in workbook:
                sheet.reset_dimensions()
                numbers=_numeric_cells(archive,sheet.title)
                sheets[sheet.title]=[(i,{j:_cell_value(c,numbers) for j,c in enumerate(row,1)})
                                     for i,row in enumerate(sheet.iter_rows(),1)]
                uncached[sheet.title]={coordinate for coordinate,cell in numbers.items()
                                      if cell.formula and not cell.token}
                element=next(s for s in document.findall('s:sheets/s:sheet',N) if s.attrib['name']==sheet.title)
                target=next(r.attrib['Target'] for r in links if r.attrib['Id']==element.attrib[f'{{{REL}}}id'])
                target=target.lstrip('/') if target.startswith('/') else normpath(join('xl',target))
                root=ET.fromstring(archive.read(target))
                merged[sheet.title]={}
                for cell in root.findall('s:mergeCells/s:mergeCell',N):
                    start,row,end,last=range_boundaries(cell.attrib['ref'])
                    if row==last:
                        merged[sheet.title][row,start]=end
        return WorkbookData(str(path.resolve()),sha256(payload).hexdigest(),sheets,uncached,merged)
    except Exception as exc:
        if isinstance(exc,ExcelImportError):
            raise
        raise ExcelImportError('Не удалось прочитать Excel-файл') from exc
    finally:
        if workbook is not None:
            workbook.close()
