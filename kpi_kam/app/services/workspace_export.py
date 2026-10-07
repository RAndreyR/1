"""Administrative XLSX export of persisted engine values, without recalculation."""
from decimal import Decimal, ROUND_HALF_UP, localcontext
from pathlib import Path
from tempfile import NamedTemporaryFile
import os
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
from app.services.sales_importer import client_identity
from app.services.snapshot_service import load_object

SUMMARY_HEADERS=['ФИО','роль','год','квартал','план','positive sales','returns','actual after returns',
    '% выполнения','eligible base B1','eligible base B2','eligible base B3','eligible base B4',
    'calls plan','calls fact','calls bonus','calculated premium','return clawback','carried clawback',
    'carried in','payable','paid','status','price version','calculation ID']

MONEY_FORMAT='[$-419]0.0'
PERCENT_FORMAT='[$-419]0.0%'
MONEY_HEADERS={'план','positive sales','returns','actual after returns',
    *(f'eligible base B{i}' for i in range(1,5)),
    'calls bonus','calculated premium','return clawback','carried clawback','carried in','payable','paid',
    'Выручка','Цена 1 знак','Порог 1 знак','Сумма','Остаток суммы','Clawback','План отгрузок'}
PERCENT_HEADERS={'% выполнения','Ставка','Исходная ставка'}


def _rounded(value,step):
    with localcontext() as context:
        context.prec=max(50,len(value.as_tuple().digits)+abs(value.as_tuple().exponent)+4)
        return value.quantize(Decimal(step),rounding=ROUND_HALF_UP)


def _row(sheet,values):
    # Decimal stays numeric in XLSX. Strings from source files stay literal text.
    sheet.append(values)
    for i,cell in enumerate(sheet[sheet.max_row],1):
        header=sheet.cell(1,i).value
        number_format=MONEY_FORMAT if header in MONEY_HEADERS else PERCENT_FORMAT if header in PERCENT_HEADERS else None
        if sheet.title=='Параметры' and header=='Значение' and len(values)>=8:
            parameter=values[6]
            if parameter in ('plan gate','calls gate') or parameter in {f'B{j} rate' for j in range(1,5)}:
                number_format=PERCENT_FORMAT
            elif parameter=='calls max':number_format=MONEY_FORMAT
        if number_format:
            cell.number_format=number_format
            if isinstance(cell.value,(Decimal,int)) and not isinstance(cell.value,bool):
                cell.value=_rounded(Decimal(cell.value),'0.001' if number_format==PERCENT_FORMAT else '0.1')
        if isinstance(cell.value,str):cell.data_type='s'


def export_workspace(repository,path,year,employee_id=None,quarter=None):
    if type(year) is not int or not 1900<=year<=2100 or quarter not in (None,1,2,3,4):
        raise ValueError('Выберите год и квартал Q1–Q4 либо весь год')
    path=Path(path).resolve()
    if path.suffix.casefold()!='.xlsx':raise ValueError('Экспорт должен иметь расширение .xlsx')
    sources=[r[0] for r in repository.connection.execute('SELECT source_file FROM sales_imports UNION SELECT source_file FROM price_list_versions UNION SELECT source_file FROM imports')]
    if path==Path(repository.path).resolve() or any(path==Path(p).resolve() for p in sources):
        raise ValueError('Нельзя перезаписывать исходный Excel или базу данных')
    book=Workbook();book.remove(book.active)
    headers={
        'Сводка':SUMMARY_HEADERS,
        'Включенные отгрузки':['ФИО','Год','Месяц','Квартал','Клиент','Тип','Исходный продукт','Канонический продукт','Группа','Количество','Выручка','Цена 1 знак','Порог 1 знак','Блок','Ставка','Лист','Строка','Контракт','Event ID','Price version'],
        'Исключенные отгрузки':['ФИО','Год','Месяц','Квартал','Клиент','Тип','Исходный продукт','Канонический продукт','Группа','Количество','Выручка','Цена 1 знак','Порог 1 знак','Блок','Ставка','Лист','Строка','Контракт','Event ID','Причина','Price version'],
        'Возвраты':['ФИО','Год','Месяц','Квартал','Клиент','Продукт','Количество','Сумма','Остаток количества','Остаток суммы','Clawback','Статус','Лист','Строка','Контракт','Return ID'],
        'Корректировки':['ФИО','Return ID','Shipment ID','Продукт','Группа','Тип','Количество','Сумма','Eligible','Блок','Исходная ставка','Исходная роль','Исходный год','Исходный квартал','Paid','Clawback','Год корректировки','Квартал корректировки'],
        'Планы':['ФИО','Год','Роль','Квартал','План отгрузок','План звонков','Факт звонков'],
        'Параметры':['Тип','ФИО','Год','Расчет','Price version','Источник','SHA256 / Параметр','Значение']}
    for name,values in headers.items():_row(book.create_sheet(name),values)
    for name in ('Включенные отгрузки','Исключенные отгрузки','Возвраты'):
        book[name].cell(1,len(headers[name])+1,'Фасовка')
    snapshots={}
    for entry in repository.workspace_history(employee_id,year):
        snapshot=repository.workspace_snapshot(entry[0])
        snapshots.setdefault(snapshot.employee.id,snapshot)
    employees={e.id:e for e in repository.employees(True)}
    selected_quarters=(quarter,) if quarter else (1,2,3,4)
    for eid,snapshot in snapshots.items():
        calc=snapshot.calculation;name=employees[eid].name
        payments={p.quarter:p for p in repository.payments(eid,year)}
        for q in selected_quarters:
            summary=calc.quarters[q-1];result=calc.result.quarters[q-1];payment=payments.get(q)
            _row(book['Сводка'],[name,payment.snapshot.profile.role if payment else snapshot.profile.role,year,q,result.plan,
                summary.positive_sales,summary.returns,summary.actual,
                result.achievement if result.achievement is not None else 'Не определено',
                *(b.base for b in result.blocks),summary.calls_plan,summary.calls_fact,summary.calls_bonus,
                summary.calculated_premium,summary.clawback,summary.carried_out,summary.carried_in,summary.payable,
                payment.amount if payment else summary.paid,'paid/closed' if payment else summary.status,
                payment.snapshot.price.id if payment else snapshot.price.id,snapshot.id])
        for event in calc.event_audit:
            e,a=event.event,event.audit
            if e.quarter not in selected_quarters:continue
            values=[name,e.year,e.month,e.quarter,a.shipment.client,a.client_type,e.product_raw,
                a.product.canonical_name if a.product else '',a.product.category if a.product else '',
                e.quantity,e.revenue,a.actual_price_rounded,a.threshold_rounded,a.block,event.rate,e.source_sheet,e.source_row,e.contract,e.event_id]
            _row(book['Включенные отгрузки'] if a.eligible else book['Исключенные отгрузки'],
                 values+([] if a.eligible else [a.exclusion_reason])+[payments[e.quarter].snapshot.price.id if e.quarter in payments else snapshot.price.id,e.packaging])
        for review in calc.return_reviews:
            e=review.event
            if e.quarter in selected_quarters:
                _row(book['Возвраты'],[name,e.year,e.month,e.quarter,client_identity(e.db,e.lpu)[1],e.product_raw,
                    e.quantity,e.revenue,review.remaining_quantity,review.remaining_revenue,review.clawback,review.status,
                    e.source_sheet,e.source_row,e.contract,e.event_id,e.packaging])
        for a in calc.allocations:
            if a.target_year==year and a.target_quarter in selected_quarters:
                _row(book['Корректировки'],[name,a.return_id,a.shipment_id,a.product.canonical_name if a.product else '',
                    a.category,a.client_type,a.quantity,a.revenue,a.eligible,a.block,a.rate,a.original_role,
                    a.original_year,a.original_quarter,a.paid,a.clawback,a.target_year,a.target_quarter])
        _row(book['Параметры'],['Прайс snapshot',name,year,snapshot.id,snapshot.price.id,snapshot.price.source_file,snapshot.price.source_hash,snapshot.price.imported_at])
        for label,value in [('plan gate',snapshot.policy.rules.plan_gate),*[(f'B{i+1} rate',r) for i,r in enumerate(snapshot.policy.rules.rates)],
                            ('calls gate',snapshot.policy.calls_gate),('calls max',snapshot.policy.calls_max),('Q4 formula','Annual calculated premium - actually paid Q1-Q3, then current clawback')]:
            _row(book['Параметры'],['Policy snapshot',name,year,snapshot.id,snapshot.price.id,'',label,value])
    for eid,profile_year,role,plans_json,calls_plans,calls_facts in repository.connection.execute('SELECT employee_id,year,role,plans_json,calls_plans_json,calls_facts_json FROM employee_year_profiles WHERE year=?',(year,)):
        if employee_id is not None and eid!=employee_id:continue
        plans=load_object(plans_json) if plans_json else None
        calls_plans,calls_facts=load_object(calls_plans),load_object(calls_facts)
        for q in selected_quarters:
            _row(book['Планы'],[employees[eid].name,profile_year,role,q,plans.values[q-1] if plans else 'Не задан',calls_plans[q-1],calls_facts[q-1]])
    for sheet in book:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='165DCE')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width=min(55,max(14,max(len(str(c.value or '')) for c in column)+2))
    path.parent.mkdir(parents=True,exist_ok=True)
    temporary=None
    try:
        with NamedTemporaryFile(dir=path.parent,suffix='.xlsx',delete=False) as stream:temporary=Path(stream.name)
        book.save(temporary);os.replace(temporary,path)
    finally:
        book.close()
        if temporary is not None and temporary.exists():temporary.unlink()
    return path
