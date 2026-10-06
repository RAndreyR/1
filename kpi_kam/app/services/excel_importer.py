"""Read XLSX/XLSM without executing formulas/macros or changing the source.

openpyxl decodes sheets/text/types. Original XML numeric tokens are converted
straight to Decimal to preserve financial precision, including cached formulas.
"""
from collections.abc import Iterator, Mapping
from dataclasses import dataclass
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import logging
from pathlib import Path
from posixpath import join, normpath
import re
from uuid import NAMESPACE_URL, uuid5
from xml.etree import ElementTree as ET
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from openpyxl.cell.read_only import ReadOnlyCell, EmptyCell
from openpyxl.utils.exceptions import InvalidFileException

from app.models.domain import Product, Shipment
from app.models.import_data import ExcelImport, SheetLayout, ValidationIssue
from app.utils.normalization import CalculationInputError, decimal_value, normalize_text

logger = logging.getLogger(__name__)
N = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

SHIPMENT_HEADERS = {
    'quarter': ('Период',), 'client': ('Клиент',),
    'client_type': ('Тип клиента',), 'product_raw': ('Номенклатура',),
    'quantity': ('Количество',), 'revenue': ('Выручка',),
    'unit_price': ('цена за 1 кг',), 'legal_entity': ('ЮЛ',),
    'bitrix_task': ('Задача в Битрикс',), 'comment': ('Комментарии',),
}
PRICE_HEADERS = {
    'canonical_name': ('Продукт', 'Номенклатура', 'Наименование продукта', 'Товар'),
    'category': ('Тип продукта', 'Категория', 'Категория продукта'),
    'price_lpu': ('зеленая зона ЛПУ', 'зеленая зона для ЛПУ'),
    'price_distributor': ('зеленая зона для дистров', 'зеленая зона дистрибьютер',
                          'зеленая зона для дистрибьюторов', 'зеленая зона для дистрибьютеров'),
}


class ExcelImportError(ValueError):
    """The file cannot be opened as a supported workbook."""


def canonical_product_id(name: str) -> str:
    """Stable across price-list row moves and whitespace/case changes."""
    return uuid5(NAMESPACE_URL, 'kpi-kam:product:' + normalize_text(name)).hex


def _header(value: object) -> str:
    return normalize_text(value).replace('ё', 'е') if isinstance(value, str) else ''


@dataclass(frozen=True)
class NumericCell:
    token: str | None
    formula: bool
    numeric: bool


def _numeric_cells(archive: ZipFile, sheet_name: str) -> dict[str, NumericCell]:
    workbook = ET.fromstring(archive.read('xl/workbook.xml'))
    sheet = next(s for s in workbook.findall('s:sheets/s:sheet', N) if s.attrib['name'] == sheet_name)
    links = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
    target = next(r.attrib['Target'] for r in links if r.attrib['Id'] == sheet.attrib[f'{{{REL}}}id'])
    target = target.lstrip('/') if target.startswith('/') else normpath(join('xl', target))
    result = {}
    with archive.open(target) as stream:
        for _, element in ET.iterparse(stream, events=('end',)):
            if element.tag == f"{{{N['s']}}}c":
                kind = element.attrib.get('t', 'n')
                formula = element.find('s:f', N) is not None
                if kind == 'n' or formula:
                    result[element.attrib['r']] = NumericCell(
                        element.findtext('s:v', namespaces=N), formula, kind == 'n')
                element.clear()
            elif element.tag == f"{{{N['s']}}}row":
                element.clear()
    return result


def _cell_value(cell: ReadOnlyCell | EmptyCell, numbers: Mapping[str, NumericCell]) -> object:
    entry = numbers.get(getattr(cell, 'coordinate', ''))
    if entry and entry.numeric and cell.data_type != 'd':
        # A formatted Excel date is not a monetary number.
        if not entry.token:
            return None
        try:
            return decimal_value(entry.token)
        except CalculationInputError:
            return entry.token  # Field parser will report the cell's location.
    return cell.value


def _parse_number(value: object, required: bool) -> Decimal | None:
    if value is None or isinstance(value, str) and not value.strip():
        if required:
            raise CalculationInputError('Обязательное числовое значение отсутствует')
        return None
    if isinstance(value, str):
        text = value.strip()
        # Allow Russian decimal commas and correctly grouped thousands only.
        if any(c.isspace() for c in text):
            text = re.sub(r'\s', ' ', text)
            if not re.fullmatch(r'[+-]?\d{1,3}(?: \d{3})+(?:[.,]\d+)?', text):
                raise CalculationInputError('Некорректная группировка числа')
            text = text.replace(' ', '')
        if not re.fullmatch(r'[+-]?(?:\d+(?:[.,]\d*)?|[.,]\d+)(?:[eE][+-]?\d+)?', text):
            raise CalculationInputError('Некорректное числовое значение')
        value = text.replace(',', '.')
    return decimal_value(value)


def _text(value: object) -> str | None:
    if value is None:
        return None
    return str(value)


def _layout(rows: list[tuple[int, dict[int, object]]], name: str,
            headers: Mapping[str, tuple[str, ...]], issues: list[ValidationIssue]) -> SheetLayout | None:
    known = {_header(alias): field for field, aliases in headers.items() for alias in aliases}
    candidates = []
    for number, cells in rows:
        found: dict[str, list[int]] = {}
        for column, value in cells.items():
            field = known.get(_header(value))
            if field:
                found.setdefault(field, []).append(column)
        if len(found) >= 2:
            candidates.append((number, found))
    if not candidates:
        issues.append(ValidationIssue('header_not_found', 'Не найдена строка заголовков', 'error', name))
        return None
    number, found = max(candidates, key=lambda item: len(item[1]))
    for field in (field for field in headers if field not in found):
        issues.append(ValidationIssue('missing_column', f'Отсутствует столбец: {headers[field][0]}',
                                      'error', name, number, field))
    for field, columns in found.items():
        if len(columns) > 1:
            issues.append(ValidationIssue('duplicate_column', f'Повторяется столбец: {headers[field][0]}',
                                          'error', name, number, field))
    return SheetLayout(name, number, {field: cols[0] for field, cols in found.items()})


def _data_rows(
    rows: list[tuple[int, dict[int, object]]], layout: SheetLayout,
) -> Iterator[tuple[int, dict[str, object]]]:
    for number, cells in rows:
        if number <= layout.header_row:
            continue
        values = {field: cells.get(column) for field, column in layout.columns.items()}
        if any(value is not None and (not isinstance(value, str) or value.strip()) for value in values.values()):
            yield number, values


def import_excel(path: str | Path) -> ExcelImport:
    """Return parsed data and structural errors; never silently drop invalid rows.

    Errors block final calculation via ImportSession. Plans are never imported.
    Business exclusions (unknown products/clients/quarters) remain in shipments.
    """
    path = Path(path)
    if path.suffix.casefold() not in ('.xlsx', '.xlsm'):
        raise ExcelImportError('Поддерживаются только файлы .xlsx и .xlsm')
    logger.info('Excel import started')
    try:
        payload = path.read_bytes()
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True,
                                 keep_vba=False, keep_links=False)
    except (OSError, BadZipFile, InvalidFileException, KeyError, ET.ParseError, ValueError) as exc:
        logger.warning('Excel file could not be opened: %s', type(exc).__name__)
        raise ExcelImportError('Не удалось открыть Excel-файл') from exc
    issues: list[ValidationIssue] = []
    layouts = []
    shipments, products = [], []
    counts = {'отгрузки': 0, 'прайс': 0}
    detected_sheets = tuple(workbook.sheetnames)
    try:
        with ZipFile(BytesIO(payload)) as archive:
            for logical, headers in (('отгрузки', SHIPMENT_HEADERS), ('прайс', PRICE_HEADERS)):
                matches = [name for name in detected_sheets if normalize_text(name) == logical]
                if len(matches) != 1:
                    code = 'missing_sheet' if not matches else 'ambiguous_sheet'
                    issues.append(ValidationIssue(code, f'Не найден однозначный лист «{logical}»', 'error', logical))
                    continue
                name = matches[0]
                sheet = workbook[name]
                sheet.reset_dimensions()  # Trust actual cells, not stale worksheet dimensions.
                numbers = _numeric_cells(archive, name)
                rows = [(number, {col: _cell_value(cell, numbers) for col, cell in enumerate(row, 1)})
                        for number, row in enumerate(sheet.iter_rows(), 1)]
                layout = _layout(rows, name, headers, issues)
                if layout is None:
                    continue
                layouts.append(layout)
                seen_products = set()
                for number, values in _data_rows(rows, layout):
                    counts[logical] += 1
                    # Incomplete/ambiguous headers are never used to parse partial data.
                    if any(i.severity == 'error' and i.sheet == name and i.source_row == layout.header_row for i in issues):
                        continue
                    parsed = {}
                    failed = False
                    numeric_fields = ('quantity', 'revenue', 'unit_price') if logical == 'отгрузки' else ('price_lpu', 'price_distributor')
                    for field in numeric_fields:
                        try:
                            parsed[field] = _parse_number(values[field], field in ('quantity', 'revenue'))
                        except CalculationInputError as exc:
                            issues.append(ValidationIssue('invalid_number', str(exc), 'error', name, number, field))
                            failed = True
                        coordinate = f'{_column_letter(layout.columns[field])}{number}'
                        cell = numbers.get(coordinate)
                        if cell and cell.formula and not cell.token:
                            issues.append(ValidationIssue('uncached_formula', 'У формулы нет сохраненного результата',
                                'error' if field in ('quantity', 'revenue') else 'warning', name, number, field))
                    if failed:
                        continue
                    if logical == 'отгрузки':
                        shipments.append(Shipment(number, _text(values['quarter']), _text(values['client_type']),
                            _text(values['product_raw']), parsed['quantity'], parsed['revenue'], parsed['unit_price'],
                            _text(values['client']) or '', _text(values['legal_entity']) or '',
                            _text(values['bitrix_task']) or '', _text(values['comment']) or ''))
                    else:
                        raw_name, raw_category = values['canonical_name'], values['category']
                        if not isinstance(raw_name, str) or not normalize_text(raw_name) or not isinstance(raw_category, str):
                            issues.append(ValidationIssue('invalid_product', 'Не заданы название и категория продукта',
                                                          'error', name, number))
                            continue
                        normalized = normalize_text(raw_name)
                        if normalized in seen_products:
                            issues.append(ValidationIssue('duplicate_product', 'Повторяется нормализованное название продукта',
                                                          'error', name, number, 'canonical_name'))
                            continue
                        seen_products.add(normalized)
                        try:
                            products.append(Product(canonical_product_id(raw_name), raw_name, raw_category,
                                                    parsed['price_lpu'], parsed['price_distributor']))
                        except CalculationInputError as exc:
                            issues.append(ValidationIssue('invalid_product', str(exc), 'error', name, number))
                if counts[logical] == 0:
                    issues.append(ValidationIssue('empty_sheet', 'Лист не содержит строк данных', 'error', name))
    except (BadZipFile, KeyError, ET.ParseError, ValueError, OSError) as exc:
        logger.warning('Excel content could not be read: %s', type(exc).__name__)
        raise ExcelImportError('Повреждена структура Excel-файла') from exc
    finally:
        workbook.close()
    for issue in issues:
        logger.warning('Excel validation: code=%s sheet=%s row=%s field=%s',
                       issue.code, issue.sheet, issue.source_row, issue.field)
    logger.info('Excel import completed: shipments=%d products=%d errors=%d',
                len(shipments), len(products), sum(i.severity == 'error' for i in issues))
    return ExcelImport(str(path.resolve()), sha256(payload).hexdigest(), detected_sheets,
                       tuple(layouts), tuple(shipments), tuple(products), tuple(issues),
                       counts['отгрузки'], counts['прайс'])


def _column_letter(index: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(index)
