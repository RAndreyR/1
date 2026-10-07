"""Verify frozen Qt/plugins/Excel/Decimal/SQLite in the actual built process.

Input fixture and frozen expectations are supplied externally by CI, not bundled.
No business formulas or alternate calculation engine live here.
"""
from decimal import Decimal
import json
import os
from pathlib import Path
import platform
import tempfile
import traceback


def run_packaging_smoke(report_path: str, fixture_path: str, expected_path: str) -> int:
    os.environ['QT_QPA_PLATFORM'] = 'offscreen'
    report = Path(report_path).resolve()
    report.parent.mkdir(parents=True, exist_ok=True)
    status = {'status':'failed', 'platform':platform.system(), 'checks':[]}
    window = None
    try:
        from PySide6.QtCore import QElapsedTimer, QTimer
        from PySide6.QtWidgets import QApplication
        from app.ui.main_window import MainWindow
        from app.utils.normalization import format_money

        expected = json.loads(Path(expected_path).read_text(encoding='utf-8'))
        fixture = Path(fixture_path).resolve()
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory(prefix='kpi-kam-build-') as directory:
            window = MainWindow(Path(directory)/'kpi_kam.db')
            window.new_page.employee.setText('Проверка сборки')
            window.new_page.year.setValue(2026)
            # Explicit regression-test plans; never read plans from the workbook.
            for field,value in zip(window.new_page.plan_inputs,('20000000','10000000','0','0')):
                field.setText(value)
            window.new_page.source.setText(str(fixture))
            window.show()
            timer = QElapsedTimer()
            timer.start()
            failures = []

            def finish_checks() -> None:
                if window._worker is not None:
                    if timer.elapsed() < 60_000:
                        QTimer.singleShot(25,finish_checks)
                        return
                    failures.append('Excel import did not finish within 60 seconds')
                    app.quit()
                    return
                try:
                    session = window.service.session
                    if session is None or session.imported.errors:
                        raise AssertionError(window.error_label.text() or 'Packaged Excel import failed')
                    if session.imported.source_hash != expected['fixture_sha256']:
                        raise AssertionError('Unexpected fixture hash')
                    if len(session.imported.shipments) != expected['shipments'] or len(session.imported.products) != expected['products']:
                        raise AssertionError('Unexpected imported row/product counts')
                    window.navigate(2)
                    for entry in session.mappings:
                        if entry.status == 'unresolved':
                            row = next(i for i,e in enumerate(session.mappings) if e.key == entry.key)
                            window.mapping_page.grid.selectRow(row)
                            window.mapping_page.leave_button.click()
                    if not window.mapping_page.calculate_button.isEnabled():
                        raise AssertionError('Packaged calculation button is not enabled')
                    window.mapping_page.calculate_button.click()
                    if window.saved is None:
                        raise AssertionError(window.error_label.text() or 'Packaged calculation failed')
                    result = window.saved.result
                    for quarter in result.quarters:
                        key = str(quarter.quarter)
                        if quarter.actual != Decimal(expected['quarter_actuals'][key]):
                            raise AssertionError(f'Actual regression mismatch Q{key}')
                        if quarter.calculated_premium != Decimal(expected['quarter_premiums'][key]):
                            raise AssertionError(f'Premium regression mismatch Q{key}')
                        if [b.base for b in quarter.blocks] != [Decimal(v) for v in expected['block_bases'][key]]:
                            raise AssertionError(f'Block regression mismatch Q{key}')
                    if len(result.included_rows) != expected['included_count']:
                        raise AssertionError('Included row count mismatch')
                    for row,reference in zip(result.audit,expected['audit'],strict=True):
                        if (row.shipment.source_row,row.eligible,row.exclusion_reason) != (reference['source_row'],reference['eligible'],reference['reason']):
                            raise AssertionError('Packaged source-row audit mismatch')
                    if window.repository.load_calculation(window.saved.id).result != result:
                        raise AssertionError('SQLite snapshot did not round-trip exactly')
                    window.open_history(window.saved.id)
                    window.show_premium(4)
                    if window.premium_page.cards['payable'].value.text() != format_money(result.quarters[3].payable):
                        raise AssertionError('Packaged premium display mismatch')
                    window.show_audit(result.quarters[0].blocks[0].source_rows,'Packaging verification')
                    if window.audit_page.selection.included_revenue != result.quarters[0].blocks[0].base:
                        raise AssertionError('Packaged drill-down mismatch')
                    if window.grab().isNull():
                        raise AssertionError('Qt rendering failed')
                    status.update(status='passed',checks=['Qt window/rendering','background Excel import',
                        'explicit product mapping','Decimal workbook regression','audit drill-down','SQLite history'],
                        shipment_count=len(result.audit),included_count=len(result.included_rows),
                        annual_actual=str(result.annual_actual),annual_premium=str(result.annual_calculated_premium))
                except Exception:
                    failures.append(traceback.format_exc())
                finally:
                    app.quit()

            window.new_page.check_button.click()
            QTimer.singleShot(25,finish_checks)
            app.exec()
            window.close()
            if failures:
                raise AssertionError('\n'.join(failures))
            if status['status'] != 'passed':
                raise AssertionError('Packaging checks did not complete')
            verify_workspace(Path(directory),app)
            status['checks'].extend(['common sales and external price','KAM/SUPPORT policy',
                                    'admin permissions','paid return clawback','XLSX export','workspace restart'])
        report.write_text(json.dumps(status,ensure_ascii=False,indent=2),encoding='utf-8')
        return 0
    except Exception:
        if window is not None:
            window.close()
        status['status'] = 'failed'
        status['error'] = traceback.format_exc()
        report.write_text(json.dumps(status,ensure_ascii=False,indent=2),encoding='utf-8')
        return 1


def verify_workspace(directory,app):
    """Check new modules inside the frozen process using synthetic workbooks."""
    from datetime import datetime
    from openpyxl import Workbook,load_workbook
    from app.models.domain import Plans
    from app.ui.workspace_window import WorkspaceWindow
    from app.services.sales_importer import FIELDS,SHEETS
    from app.services.price_importer import import_price
    from app.services.sales_importer import import_sales
    from app.services.admin_auth import DEFAULT_PASSWORD
    price_path=directory/'price.xlsx';sales_path=directory/'sales.xlsx'
    book=Workbook();sheet=book.active
    sheet.append(['Продукт','Группа','Зеленая зона ЛПУ','Зеленая зона дистрибьюторов'])
    sheet.append(['Синтетический товар 200 мл','ЭП','100','100'])
    sheet.append(['Синтетический товар 500 мл','ЭП','200','200']);book.save(price_path);book.close()

    def sales(include_return):
        book=Workbook();book.remove(book.active)
        for i,name in enumerate(SHEETS):
            sheet=book.create_sheet(name)
            for col,key in enumerate(FIELDS,1):sheet.cell(6,col,FIELDS[key])
            sheet.cell(5,10,datetime(2026,1,1));sheet.merge_cells('J5:K5')
            sheet.cell(6,10,'количество');sheet.cell(6,11,'сумма')
            sheet.cell(5,13,'Возврат июль');sheet.cell(5,14,'Возврат июль')
            sheet.append(['ФО','Регион','Трофимов','ЛПУ','Дистрибьютер','ЮЛ','С-1','Синтетический товар','200 мл',
                          '10' if i==0 else '0','1000' if i==0 else '0',None,
                          '2' if include_return and i==0 else '0','200' if include_return and i==0 else '0'])
        book.save(sales_path);book.close()

    database=directory/'workspace.db';window=WorkspaceWindow(database)
    try:
        window.show();app.processEvents()
        employee=window.repository.resolve_employee('Трофимов');flow=window.workflow
        window.employee.setCurrentIndex(window.employee.findData(employee.id))
        window.year.setValue(2026)
        assert not flow.is_admin
        assert flow.login_admin(DEFAULT_PASSWORD)
        flow.set_role(employee.id,2026,'SUPPORT')
        window.repository.save_price(*import_price(price_path))
        sales(False);window.repository.save_sales(import_sales(sales_path,2026))
        price=window.repository.price_version()
        flow.map_product('Синтетический товар',price.products[0].id,packaging='200 мл')
        flow.save_inputs(employee.id,2026,Plans('1000','1000','1000','1000'),
            calls_plans=('900','800','600','1000'),calls_facts=('840','0','0','0'))
        snapshot=flow.calculate(employee.id,2026)
        assert snapshot.calculation.event_audit[0].event.packaging=='200 мл'
        assert snapshot.calculation.event_audit[0].audit.product.id==price.products[0].id
        assert snapshot.calculation.quarters[0].calls_bonus==Decimal('28000')
        assert snapshot.calculation.quarters[0].payable==Decimal('28050')
        flow.confirm_payment(snapshot.id,1)
        sales(True);window.repository.save_sales(import_sales(sales_path,2026))
        later=flow.calculate(employee.id,2026)
        assert later.calculation.quarters[2].actual==Decimal('-200')
        assert later.calculation.quarters[2].clawback==Decimal('10')
        window.display_workspace(later)
        assert window.returns.grid.rowCount()==1
        flow.export_data(directory/'export.xlsx',2026,employee_id=employee.id)
        exported=load_workbook(directory/'export.xlsx')
        try:
            assert 'Возвраты' in exported.sheetnames
            headers={c.value:c.column for c in exported['Сводка'][1]}
            amount=exported['Сводка'].cell(2,headers['payable'])
            assert amount.data_type=='n' and amount.value==28050
            assert amount.number_format=='[$-419]0.0'
            achievement=exported['Сводка'].cell(2,headers['% выполнения'])
            assert achievement.data_type=='n' and achievement.number_format=='[$-419]0.0%'
        finally:exported.close()
        window.refresh_workspace()
        assert window.calls_plan_inputs[0].text()=='900'
        assert window.calls_plan_inputs[0].isReadOnly()
        window.mappings.grid.selectRow(0)
        assert window.mappings.products.count()==2
        assert '200 мл' in window.mappings.products.itemText(1)
        assert not window.grab().isNull()
    finally:window.close()
    reopened=WorkspaceWindow(database)
    try:
        assert not reopened.workflow.is_admin
        assert reopened.repository.payments(employee.id,2026)[0].amount==Decimal('28050')
        assert reopened.repository.workspace_snapshot(snapshot.id).status=='paid/closed'
        assert reopened.calls_plan_inputs[0].text()=='900'
    finally:reopened.close()
