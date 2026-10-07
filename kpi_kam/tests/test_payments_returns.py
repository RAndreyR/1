from dataclasses import replace
from decimal import Decimal
import sqlite3
import pytest
from app.models.domain import Plans
from app.repositories.workspace_repository import WorkspaceRepository
from app.services.workspace_service import WorkspaceService
from app.utils.normalization import CalculationInputError
from tests.event_helpers import sales_book,price_book


def ready(tmp_path,*,threshold='100'):
    repo=WorkspaceRepository(tmp_path/'app.db');service=WorkspaceService(repo)
    service.login_admin('stopp')
    employee=repo.resolve_employee('Трофимов')
    service.set_role(employee.id,2026,'KAM')
    service.import_price(price_book(tmp_path/'price.xlsx',threshold=threshold))
    service.import_sales(sales_book(tmp_path/'sales.xlsx'),2026)
    service.save_inputs(employee.id,2026,Plans('1000','1000','1000','1000'))
    return repo,service,employee


def test_paid_eligible_return_uses_original_rate_and_price(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        first=service.calculate(employee.id,2026)
        service.confirm_payment(first.id,1)
        service.import_price(price_book(tmp_path/'new-price.xlsx',category='ЭП',threshold='500'))
        service.import_sales(sales_book(tmp_path/'new-sales.xlsx',returns={7:('2','200')}),2026)
        second=service.calculate(employee.id,2026)
        a=second.calculation.allocations[0]
        assert a.paid and a.eligible
        assert a.rate==Decimal('0.05') and a.category=='ВМК'
        assert a.clawback==10
        assert second.calculation.result.green_totals[1,'дистрибьютер','ВМК']==1000
        assert second.calculation.quarters[2].actual==-200
        assert second.calculation.quarters[2].carried_out==10
        assert repo.workspace_snapshot(first.id).calculation==first.calculation
        with pytest.raises(sqlite3.IntegrityError):
            repo.connection.execute("UPDATE calculations SET workspace_snapshot_json='{}' WHERE id=?",(first.id,))
        repo.connection.rollback()


def test_manual_calls_plans_are_saved_and_paid_quarter_plan_is_protected(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        service.set_role(employee.id,2026,'SUPPORT')
        service.logout_admin()
        plans=Plans('1000','1000','1000','1000')
        profile=service.save_inputs(employee.id,2026,plans,calls_plans=('1000','800','600','900'),calls_facts=('900','0','0','0'))
        assert profile.calls_plans[0]==1000
        snapshot=service.calculate(employee.id,2026)
        assert snapshot.calculation.quarters[0].calls_bonus==27000
        service.confirm_payment(snapshot.id,1)
        with pytest.raises(CalculationInputError,match='План звонков оплаченного квартала'):
            service.save_inputs(employee.id,2026,plans,calls_plans=('999','800','600','900'))
        assert repo.year_profile(employee.id,2026)==profile
        changed=service.save_inputs(employee.id,2026,plans,calls_plans=('1000','1200','600','900'))
        assert changed.calls_plans[1]==1200
        assert repo.payments(employee.id,2026)[0].snapshot.profile.calls_plans==profile.calls_plans


def test_ineligible_paid_return_no_clawback(tmp_path):
    repo,service,employee=ready(tmp_path,threshold='200')
    with repo:
        first=service.calculate(employee.id,2026);service.confirm_payment(first.id,1)
        service.import_sales(sales_book(tmp_path/'later.xlsx',returns={7:('2','200')}),2026)
        second=service.calculate(employee.id,2026)
        assert second.calculation.allocations[0].clawback==0
        assert second.calculation.quarters[2].actual==-200


def test_return_before_payment_not_reclassified_after_payment(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        service.import_sales(sales_book(tmp_path/'later.xlsx',returns={7:('2','200')}),2026)
        first=service.calculate(employee.id,2026)
        assert first.calculation.result.quarters[0].calculated_premium==40
        service.confirm_payment(first.id,1)
        second=service.calculate(employee.id,2026)
        assert not second.calculation.allocations[0].paid
        assert second.calculation.result.quarters[0].calculated_premium==40
        assert second.calculation.quarters[2].clawback==0
        assert second.calculation.result.green_totals[1,'дистрибьютер','ВМК']==800


def test_reimport_return_is_not_applied_twice_and_partial_payment_saved(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        first=service.calculate(employee.id,2026);service.confirm_payment(first.id,1,Decimal('35'))
        path=sales_book(tmp_path/'returns.xlsx',returns={7:('2','200')})
        for _ in range(2):
            service.import_sales(path,2026)
            latest=service.calculate(employee.id,2026)
            assert latest.calculation.quarters[2].clawback==10
        assert repo.connection.execute('SELECT count(*) FROM return_events').fetchone()[0]==1
        assert repo.connection.execute('SELECT count(*) FROM clawback_ledger').fetchone()[0]==1
        assert repo.payments(employee.id)[0].amount==35
    with WorkspaceRepository(tmp_path/'app.db') as reopened:
        assert reopened.payments(employee.id)[0].amount==35
        assert reopened.workspace_snapshot(latest.id).calculation.quarters[2].carried_out==10


def test_clawback_carries_to_next_quarter(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        service.confirm_payment(service.calculate(employee.id,2026).id,1)
        service.import_sales(sales_book(tmp_path/'later.xlsx',months={1:('10','1000'),7:('1','100')},returns={4:('8','800')}),2026)
        # Quarter gate Q3 explicitly passes, annual gate still fails.
        service.save_inputs(employee.id,2026,Plans('1000','1000','100','1000'))
        result=service.calculate(employee.id,2026).calculation
        assert result.quarters[1].clawback==40
        assert result.quarters[1].carried_out==40
        assert result.quarters[2].carried_in==40
        assert result.quarters[2].payable==0
        assert result.quarters[2].carried_out==35


def test_reset_requires_admin_two_confirmations_and_backup_preserves_config(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        first=service.calculate(employee.id,2026);service.confirm_payment(first.id,1)
        with pytest.raises(CalculationInputError):
            service.reset_history(confirmed=True)
        backup=service.reset_history(confirmed=True,confirmed_again=True)
        assert backup.is_file()
        with sqlite3.connect(backup) as before:
            assert before.execute('SELECT count(*) FROM payment_records').fetchone()[0]==1
        assert not repo.workspace_history()
        assert repo.price_version() is not None
        assert repo.resolve_employee('Трофимов').id==employee.id
        assert repo.year_profile(employee.id,2026).plans==Plans('1000','1000','1000','1000')
        service.logout_admin()
        with pytest.raises(PermissionError):
            service.reset_history(confirmed=True,confirmed_again=True)


def test_cross_year_balance_and_fifo_reserves_previously_returned_quantity(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        service.confirm_payment(service.calculate(employee.id,2026).id,1)
        service.import_sales(sales_book(tmp_path/'late2026.xlsx',returns={7:('2','200')}),2026)
        old=service.calculate(employee.id,2026)
        service.confirm_payment(old.id,3)
        service.set_role(employee.id,2027,'KAM')
        service.save_inputs(employee.id,2027,Plans('1000','1000','1000','1000'))
        service.import_sales(sales_book(tmp_path/'sales2027.xlsx',year=2027,returns={2:('9','900')}),2027)
        new=service.calculate(employee.id,2027)
        assert new.calculation.quarters[0].carried_in==10
        parts=new.calculation.allocations
        assert [(a.original_year,a.quantity,a.paid) for a in parts]==[(2026,Decimal(8),True),(2027,Decimal(1),False)]
        assert not new.calculation.unallocated_returns


def test_unknown_manager_requires_admin_mapping_without_loss(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        imported=service.import_sales(sales_book(tmp_path/'unknown.xlsx',manager='Бабенкова'),2026)
        assert len(imported.shipments)==1
        assert service.unknown_managers(imported)==('Бабенкова',)
        with pytest.raises(CalculationInputError):
            service.calculate(employee.id,2026)
        new=service.add_employee('Бабенкова Ирина')
        service.map_employee('Бабенкова',new.id)
        service.set_role(new.id,2026,'SUPPORT')
        service.save_inputs(new.id,2026,Plans('0','0','0','0'))
        assert service.calculate(new.id,2026).calculation.result.annual_actual==1000


def test_unallocated_return_manual_resolution_and_admin_guard(tmp_path):
    from openpyxl import load_workbook
    from app.services.sales_importer import FIELDS,SHEETS
    repo,service,employee=ready(tmp_path)
    with repo:
        paid=service.calculate(employee.id,2026);service.confirm_payment(paid.id,1)
        path=sales_book(tmp_path/'unallocated.xlsx',returns={7:('2','200')})
        w=load_workbook(path);s=w[SHEETS[0]]
        for col in range(1,s.max_column+1):
            s.cell(8,col,s.cell(7,col).value)
        field_cols={field:next(c.column for c in s[6] if c.value==label) for field,label in FIELDS.items()}
        s.cell(8,field_cols['contract'],'K-2')
        # Only the new contract has a return; it has no positive sale.
        s.cell(8,14,'0');s.cell(8,15,'0');s.cell(7,17,'0');s.cell(7,18,'0')
        w.save(path);w.close()
        imported=service.import_sales(path,2026)
        result=service.calculate(employee.id,2026).calculation
        assert result.unallocated_returns==(imported.returns[0].event_id,)
        assert result.quarters[2].actual==-200 and result.quarters[2].clawback==0
        service.logout_admin()
        with pytest.raises(PermissionError):
            service.reconcile_manually(imported.returns[0].event_id,imported.shipments[0].event_id,employee.id)
        service.login_admin('stopp')
        service.reconcile_manually(imported.returns[0].event_id,imported.shipments[0].event_id,employee.id)
        reconciled=service.calculate(employee.id,2026).calculation
        assert not reconciled.unallocated_returns and reconciled.quarters[2].clawback==10


def test_late_earlier_return_cannot_consume_closed_allocation_twice(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        service.confirm_payment(service.calculate(employee.id,2026).id,1)
        service.import_sales(sales_book(tmp_path/'first-return.xlsx',returns={9:('8','800')}),2026)
        service.calculate(employee.id,2026)
        service.import_sales(sales_book(tmp_path/'new-return.xlsx',returns={7:('5','500'),9:('8','800')}),2026)
        result=service.calculate(employee.id,2026).calculation
        assert sum(a.quantity for a in result.allocations)==10
        assert len(result.unallocated_returns)==1
        assert result.quarters[2].clawback==50


def test_q4_literal_rule_selected_by_user_after_prior_clawback(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        service.save_inputs(employee.id,2026,Plans('0','0','0','0'))
        service.confirm_payment(service.calculate(employee.id,2026).id,1)
        service.import_sales(sales_book(tmp_path/'later.xlsx',months={1:('10','1000'),4:('5','500')},returns={4:('2.5','250')}),2026)
        second=service.calculate(employee.id,2026)
        assert second.calculation.quarters[1].payable==Decimal('12.5')
        service.confirm_payment(second.id,2)
        annual=service.calculate(employee.id,2026)
        assert annual.calculation.result.annual_calculated_premium==75
        assert annual.calculation.result.paid_q1_q3==Decimal('62.5')
        assert annual.calculation.quarters[3].payable==Decimal('12.5')


def test_explicit_unmatched_preserves_actual_and_requires_new_version_decision(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        service.import_sales(sales_book(tmp_path/'unknown-product.xlsx',raw='Нет в прайсе'),2026)
        with pytest.raises(CalculationInputError):
            service.calculate(employee.id,2026)
        service.leave_product_unmatched('Нет в прайсе')
        result=service.calculate(employee.id,2026).calculation
        assert result.result.annual_actual==1000
        assert result.result.annual_calculated_premium==0
        assert result.result.audit[0].exclusion_reason=='Не сопоставлен продукт'
        service.import_price(price_book(tmp_path/'v2.xlsx'))
        with pytest.raises(CalculationInputError):
            service.calculate(employee.id,2026)


def test_paid_product_removed_from_new_price_keeps_original_return_basis(tmp_path):
    from dataclasses import replace
    from app.services.price_importer import import_price
    repo,service,employee=ready(tmp_path)
    with repo:
        first=service.calculate(employee.id,2026)
        service.confirm_payment(first.id,1)
        data,products=import_price(price_book(tmp_path/'new.xlsx'))
        repo.save_price(data,(replace(products[0],id='different',canonical_name='Другая номенклатура'),))
        service.import_sales(sales_book(tmp_path/'later.xlsx',returns={7:('2','200')}),2026)
        assert all(kind=='historical' for _,_,kind in service.product_mappings(employee.id,2026))
        second=service.calculate(employee.id,2026)
        assert second.calculation.quarters[2].clawback==10
        assert second.calculation.allocations[0].product==first.price.products[0]


def test_late_return_in_paid_quarter_adjusts_next_open_period(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        first=service.calculate(employee.id,2026);service.confirm_payment(first.id,1)
        service.confirm_payment(service.calculate(employee.id,2026).id,2)
        service.import_sales(sales_book(tmp_path/'late.xlsx',returns={4:('2','200')}),2026)
        result=service.calculate(employee.id,2026).calculation
        assert result.quarters[1].actual==-200
        assert result.quarters[1].clawback==0
        assert result.quarters[2].clawback==10
        assert result.allocations[0].target_quarter==3


def test_unpaid_prior_year_return_recalculates_original_base(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        original=service.calculate(employee.id,2026)
        assert original.calculation.quarters[0].calculated_premium==50
        service.set_role(employee.id,2027,'KAM')
        service.save_inputs(employee.id,2027,Plans('1000','1000','1000','1000'))
        service.import_sales(sales_book(tmp_path/'next-year.xlsx',year=2027,months={1:('0','0')},returns={1:('2','200')}),2027)
        current=service.calculate(employee.id,2027)
        assert current.calculation.quarters[0].actual==-200
        assert current.calculation.quarters[0].clawback==0
        assert not current.calculation.unallocated_returns
        assert current.calculation.allocations[0].original_year==2026
        old=service.calculate(employee.id,2026)
        assert old.calculation.quarters[0].calculated_premium==40
        assert old.calculation.quarters[0].actual==1000
        service.confirm_payment(old.id,1)
        # It was already deducted before payment: never reclassify as clawback.
        assert service.calculate(employee.id,2027).calculation.quarters[0].clawback==0
