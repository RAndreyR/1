"""Build real OOXML inputs; patch raw numeric/cache XML for precision cases."""
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
from openpyxl import Workbook
from app.services.excel_importer import SHIPMENT_HEADERS, PRICE_HEADERS

SHIP_HEADERS = [names[0] for names in SHIPMENT_HEADERS.values()]
PRICE_COLUMNS = [names[0] for names in PRICE_HEADERS.values()]
N = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def make_workbook(path: Path, *, shipments=None, products=None, ship_headers=None,
                  price_headers=None, header_row=2, reorder=False, names=('отгрузки','прайс')) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    shipments = shipments if shipments is not None else [['1 кв.', 'Клиент', 'дистрибьютер', 'Товар', 10, 1000, 100, 'ООО', 'task', 'comment']]
    products = products if products is not None else [['Товар', 'СБКС', 100, 100]]
    for title, headers, rows in [(names[0], ship_headers or SHIP_HEADERS, shipments),
                                 (names[1], price_headers or PRICE_COLUMNS, products)]:
        sheet = wb.create_sheet(title)
        for index in range(header_row-1):
            sheet.append([f'Описание {index}'])
        if reorder:
            headers = list(reversed(headers))
            rows = [list(reversed(row)) for row in rows]
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
    wb.create_sheet('KPI').append(['Планы из этого листа запрещены', 99999999])
    wb.save(path)
    wb.close()
    return path


def patch_xml(path: Path, sheet_index: int, coordinate: str, *, token=None,
              formula=None, kind='n', dimension=None) -> None:
    with ZipFile(path) as source:
        files = {name: source.read(name) for name in source.namelist()}
    name = f'xl/worksheets/sheet{sheet_index}.xml'
    root = ET.fromstring(files[name])
    cell = next(c for c in root.findall('.//s:c', N) if c.attrib['r'] == coordinate)
    cell.attrib['t'] = kind
    for child in list(cell):
        cell.remove(child)
    if formula is not None:
        ET.SubElement(cell, f"{{{N['s']}}}f").text = formula
    if token is not None:
        ET.SubElement(cell, f"{{{N['s']}}}v").text = token
    if dimension is not None:
        root.find('s:dimension', N).attrib['ref'] = dimension
    files[name] = ET.tostring(root, encoding='utf-8')
    with ZipFile(path, 'w', compression=ZIP_DEFLATED) as target:
        for name, content in files.items():
            target.writestr(name, content)
