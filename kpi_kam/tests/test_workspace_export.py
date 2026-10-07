from decimal import Decimal
import pytest
from openpyxl import load_workbook
from tests.test_payments_returns import ready
from tests.event_helpers import sales_book


def test_admin_export_values_filters_details_and_formula_text(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        first=service.calculate(employee.id,2026);service.confirm_payment(first.id,1,Decimal('40'))
        service.import_sales(sales_book(tmp_path/'update.xlsx',returns={7:('2','200')}),2026)
        service.calculate(employee.id,2026)
        output=tmp_path/'export.xlsx'
        service.export_data(output,2026,employee_id=employee.id,quarter=3)
        book=load_workbook(output,data_only=False)
        assert book.sheetnames==['Сводка','Включенные отгрузки','Исключенные отгрузки','Возвраты','Корректировки','Планы','Параметры']
        rows=list(book['Сводка'].values);values=dict(zip(rows[0],rows[1]))
        assert len(rows)==2
        assert values['ФИО']=='Трофимов Дмитрий'
        assert values['returns']=='200'
        assert values['actual after returns']=='-200'
        assert values['return clawback']=='10.00'
        assert values['carried clawback']=='10.00'
        assert book['Возвраты'].max_row==2
        assert book['Включенные отгрузки'].max_row==1
        service.export_data(tmp_path/'all.xlsx',2026)
        assert load_workbook(tmp_path/'all.xlsx')['Сводка'].max_row==5
        # Export source names as text even when Excel would treat them as formulas.
        service.update_employee(employee.id,'=HYPERLINK("bad")',True)
        service.export_data(tmp_path/'text.xlsx',2026)
        cell=load_workbook(tmp_path/'text.xlsx',data_only=False)['Сводка']['A2']
        assert cell.data_type=='s'
        service.logout_admin()
        with pytest.raises(PermissionError):service.export_data(tmp_path/'denied.xlsx',2026)
        assert not (tmp_path/'denied.xlsx').exists()


def test_export_cannot_overwrite_source(tmp_path):
    repo,service,employee=ready(tmp_path)
    with repo:
        service.calculate(employee.id,2026)
        source=tmp_path/'sales.xlsx';original=source.read_bytes()
        with pytest.raises(ValueError):service.export_data(source,2026)
        assert source.read_bytes()==original
