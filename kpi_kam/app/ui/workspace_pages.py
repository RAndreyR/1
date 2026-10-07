"""Presentation controls for monthly events. Financial values come from the engine."""
from PySide6.QtCore import Qt
from PySide6.QtWidgets import (QWidget,QVBoxLayout,QHBoxLayout,QFormLayout,QLabel,
    QPushButton,QLineEdit,QComboBox,QSpinBox,QTabWidget,QCheckBox,QDialog,
    QDialogButtonBox,QScrollArea)
from openpyxl.utils import get_column_letter
from app.models.import_data import SheetLayout
from app.services.sales_importer import FIELDS
from app.ui.common import page_layout,table,fill_table


def button(text,callback,layout):
    widget=QPushButton(text);widget.clicked.connect(callback);layout.addWidget(widget)
    return widget


class DataPage(QWidget):
    def __init__(self,title,headers,subtitle=''):
        super().__init__()
        self.body=page_layout(self,title,subtitle)
        self.search=QLineEdit();self.search.setPlaceholderText('Поиск по всем колонкам')
        self.body.addWidget(self.search)
        self.grid=table(headers);self.body.addWidget(self.grid,1)
        self.rows=[];self.keys=[];self.visible_keys=[]
        self.search.textChanged.connect(self.refresh)

    def set_rows(self,rows,keys=None):
        self.rows=rows;self.keys=list(range(len(rows))) if keys is None else list(keys)
        self.refresh()

    def refresh(self,*_):
        query=self.search.text().strip().casefold()
        indices=[i for i,row in enumerate(self.rows) if not query or query in ' '.join(str(v) for v in row).casefold()]
        self.visible_keys=[self.keys[i] for i in indices]
        fill_table(self.grid,[self.rows[i] for i in indices])

    def selected_key(self):
        row=self.grid.currentRow()
        return self.visible_keys[row] if 0<=row<len(self.visible_keys) else None


class MappingPage(DataPage):
    def __init__(self,window):
        super().__init__('Сопоставление товаров',['Из файла продаж','Фасовка','Продукт прайса','Группа','Статус','Фасовка прайса'],
            'Название и фасовка сопоставляются вместе. 200 мл — сиппинг; 500/1000 мл — реторты. Группа берется из внешнего прайса.')
        actions=QHBoxLayout();self.body.addLayout(actions)
        self.products=QComboBox();actions.addWidget(self.products,1)
        self.apply=button('Подтвердить и сохранить алиас',window.apply_product_mapping,actions)
        self.leave=button('Оставить вне зеленой зоны',window.leave_product_mapping,actions)
        self.grid.itemSelectionChanged.connect(window.refresh_mapping_products)


class AdminPanel(QWidget):
    def __init__(self,window):
        super().__init__()
        body=page_layout(self,'Администратор','Сессия действует до выхода или закрытия приложения.')
        login_row=QHBoxLayout();body.addLayout(login_row)
        self.password=QLineEdit();self.password.setEchoMode(QLineEdit.EchoMode.Password)
        self.password.setPlaceholderText('Пароль администратора');login_row.addWidget(self.password)
        self.login=button('Войти',window.admin_login,login_row)
        self.logout=button('Выйти',window.admin_logout,login_row)
        self.state=QLabel();body.addWidget(self.state)
        self.tabs=QTabWidget();body.addWidget(self.tabs,1)
        self.forms=[]
        for title in ('Прайс-лист','Сотрудники','Настройки KPI','Экспорт','История / Сброс','Смена пароля'):
            page=QWidget();layout=QVBoxLayout(page);layout.setAlignment(Qt.AlignmentFlag.AlignTop)
            scroll=QScrollArea();scroll.setWidgetResizable(True);scroll.setWidget(page)
            self.tabs.addTab(scroll,title);self.forms.append(layout)
        price=self.forms[0]
        self.price_info=QLabel();self.price_info.setWordWrap(True);price.addWidget(self.price_info)
        button('Импорт отдельного прайс-листа',window.choose_price,price)
        self.price_grid=table(['Продукт','Фасовка','Группа','ЛПУ','Дистрибьютер']);price.addWidget(self.price_grid)
        employee=self.forms[1]
        self.unknown=QLabel();self.unknown.setWordWrap(True);employee.addWidget(self.unknown)
        self.employees=QComboBox();employee.addWidget(self.employees)
        form=QFormLayout();employee.addLayout(form)
        self.full_name=QLineEdit();form.addRow('ФИО (для добавления / изменения)',self.full_name)
        self.active=QCheckBox('Активен');self.active.setChecked(True);form.addRow(self.active)
        self.alias=QLineEdit();form.addRow('Имя из файла / source alias',self.alias)
        self.year=QSpinBox();self.year.setRange(1900,2100);self.year.setValue(2026);form.addRow('Год должности',self.year)
        self.role=QComboBox();self.role.addItem('Не задана',None);self.role.addItem('КАМ','KAM');self.role.addItem('Менеджер сопровождения','SUPPORT');form.addRow('Должность',self.role)
        button('Добавить сотрудника',window.add_employee,employee)
        button('Сохранить ФИО и активность',window.edit_employee,employee)
        button('Сохранить алиас для выбранного сотрудника',window.map_employee,employee)
        self.save_role_button=button('Задать должность на год',window.save_employee_role,employee)
        self.role_feedback=QLabel();self.role_feedback.setWordWrap(True);employee.addWidget(self.role_feedback)
        self.employees.currentIndexChanged.connect(window.admin_employee_selected)
        self.year.valueChanged.connect(window.load_employee_role)
        self.role.currentIndexChanged.connect(window.employee_role_edited)
        policy=self.forms[2]
        self.policy_role=QComboBox();self.policy_role.addItems(['KAM','SUPPORT']);policy.addWidget(self.policy_role)
        self.policy_fields=[];form=QFormLayout();policy.addLayout(form)
        for label in ('Порог плана, %','Блок 1, %','Блок 2, %','Блок 3, %','Блок 4, %','Порог звонков, %','Максимум за звонки, ₽','План звонков по умолчанию'):
            edit=QLineEdit();self.policy_fields.append(edit);form.addRow(label,edit)
        self.policy_role.currentTextChanged.connect(window.load_policy_fields)
        button('Сохранить настройки для новых расчетов',window.save_role_policy,policy)
        export=self.forms[3];form=QFormLayout();export.addLayout(form)
        self.export_employee=QComboBox();form.addRow('Сотрудник',self.export_employee)
        self.export_year=QSpinBox();self.export_year.setRange(1900,2100);self.export_year.setValue(2026);form.addRow('Год',self.export_year)
        self.export_quarter=QComboBox();self.export_quarter.addItem('Весь год',None)
        for q in range(1,5):self.export_quarter.addItem(f'Q{q}',q)
        form.addRow('Период',self.export_quarter)
        button('Экспорт XLSX',window.choose_export,export)
        reset=self.forms[4]
        text=QLabel('Перед сбросом создается резервная копия БД. Сотрудники, алиасы, прайсы, планы, должности и пароль сохраняются.')
        text.setWordWrap(True);reset.addWidget(text)
        button('Обнулить историю расчетов',window.reset_history,reset)
        password=self.forms[5];form=QFormLayout();password.addLayout(form)
        self.old_password=QLineEdit();self.new_password=QLineEdit();self.repeat_password=QLineEdit()
        for label,edit in (('Текущий пароль',self.old_password),('Новый пароль',self.new_password),('Повтор нового пароля',self.repeat_password)):
            edit.setEchoMode(QLineEdit.EchoMode.Password);form.addRow(label,edit)
        button('Изменить пароль',window.change_admin_password,password)


class ColumnMappingDialog(QDialog):
    """An administrator explicitly chooses metadata; no fixed-column defaults."""
    def __init__(self,sheet_name,rows,parent=None):
        super().__init__(parent);self.sheet_name=sheet_name
        max_columns=max((max(cells,default=0) for _,cells in rows),default=0)
        rows=[tuple(cells.get(i) for i in range(1,max_columns+1)) for _,cells in rows]
        self.setWindowTitle(f'Подтвердите колонки: {sheet_name}');self.resize(1050,750)
        body=QVBoxLayout(self)
        description=QLabel('Не найдены обязательные заголовки. Выберите строку под месячными группами и колонки по предпросмотру. Укажите колонку фасовки, если она есть, и единицу для числовых значений. Данные начинаются со следующей строки.')
        description.setWordWrap(True);body.addWidget(description)
        form=QFormLayout();body.addLayout(form)
        self.header=QSpinBox();self.header.setRange(1,30);form.addRow('Строка перед данными',self.header)
        self.columns={}
        for field,label in FIELDS.items():
            combo=QComboBox();combo.addItem('Выберите колонку',None)
            for index in range(max_columns):
                samples=[]
                for row in rows[:50]:
                    value=row[index] if index<len(row) else None
                    if value is not None and str(value).strip() and str(value) not in samples:
                        samples.append(str(value))
                combo.addItem(f'{get_column_letter(index+1)}: '+ ' / '.join(samples[:3])[:120],index+1)
            form.addRow(label,combo);self.columns[field]=combo
        self.packaging_unit=QComboBox();self.packaging_unit.addItem('Единица указана в ячейке / заголовке','')
        for unit in ('мл','л','г','кг'):self.packaging_unit.addItem(unit,unit)
        form.addRow('Единица фасовки (для чисел без единицы)',self.packaging_unit)
        preview=table(['Строка']+[get_column_letter(i+1) for i in range(max_columns)])
        fill_table(preview,[[i+1,*list(row),*(['']*(max_columns-len(row)))] for i,row in enumerate(rows[:35])]);body.addWidget(preview,1)
        buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Ok|QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept);buttons.rejected.connect(self.reject);body.addWidget(buttons)

    def layout_value(self):
        return SheetLayout(self.sheet_name,self.header.value(),
            {key:combo.currentData() for key,combo in self.columns.items() if combo.currentData() is not None},
            self.packaging_unit.currentData())
