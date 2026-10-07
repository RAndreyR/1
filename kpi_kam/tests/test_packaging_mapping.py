"""Packaging from common sales must isolate aliases, thresholds and returns."""
from dataclasses import replace
from decimal import Decimal
import hashlib
import json
import pytest
from openpyxl import Workbook, load_workbook

from app.models.domain import Plans, Product
from app.models.events import ShipmentEvent, ReturnEvent, RolePolicy
from app.repositories.workspace_repository import WorkspaceRepository
from app.services.admin_auth import AdminAuth
from app.services.event_calculation import calculate_events
from app.services.product_matcher import ProductMatcher
from app.services.sales_importer import import_sales, SHEETS
from app.services.snapshot_service import dump_object, load_object
from app.services.workspace_service import WorkspaceService
from app.utils.normalization import CalculationInputError, normalize_text
from app.utils.packaging import product_mapping_key
from tests.event_helpers import sales_book


def packaged_sales(path, sizes=(200, 500, 1000), *, shift=0,returns=None,months=None):
    sales_book(path, raw='Стандарт', shift=shift,returns=returns,months=months)
    book=load_workbook(path);sheet=book[SHEETS[0]];header=6+shift
    col=next(c.column for c in sheet[header] if c.value=='Фасовка')
    sheet.cell(header,col,'Фасовка, мл')
    for row,size in enumerate(sizes,header+1):
        if row>header+1:
            for c in range(1,sheet.max_column+1):sheet.cell(row,c,sheet.cell(header+1,c).value)
        sheet.cell(row,col,size)
    book.save(path);book.close();return path


def price_variants(path):
    book=Workbook();sheet=book.active
    sheet.append(['Продукт','Группа','Зеленая зона ЛПУ','Зеленая зона дистрибьюторов'])
    for size,threshold in ((200,154),(500,308),(1000,416)):
        sheet.append([f'Иннованта Стандарт {size} мл','ЭП',threshold,threshold])
    book.save(path);book.close();return path


@pytest.mark.parametrize('shift',[0,3,12])
def test_import_reads_packaging_dynamically_and_keeps_raw_name(tmp_path,shift):
    imported=import_sales(packaged_sales(tmp_path/'sales.xlsx',shift=shift),2026)
    assert not imported.errors
    assert [(e.product_raw,e.packaging) for e in imported.shipments]==[
        ('Стандарт','200 мл'),('Стандарт','500 мл'),('Стандарт','1000 мл')]
    assert len({e.event_id for e in imported.shipments})==3


def test_packaging_aliases_are_independent_and_survive_restart(tmp_path):
    path=tmp_path/'app.db'
    with WorkspaceRepository(path) as repo:
        service=WorkspaceService(repo);assert service.login_admin('stopp')
        employee=repo.resolve_employee('Трофимов')
        service.set_role(employee.id,2026,'KAM')
        service.save_inputs(employee.id,2026,Plans('1000','1000','1000','1000'))
        price=service.import_price(price_variants(tmp_path/'price.xlsx'))
        service.import_sales(packaged_sales(tmp_path/'sales.xlsx',months={1:('10','2000')}),2026)
        assert len(service.product_mappings(employee.id,2026))==3
        for size,product in zip((200,500,1000),price.products):
            service.map_product('Стандарт',product.id,packaging=f'{size} мл')
        saved=service.calculate(employee.id,2026)
        assert [a.audit.product.id for a in saved.calculation.event_audit]==[p.id for p in price.products]
        assert [a.audit.threshold_raw for a in saved.calculation.event_audit]==list(map(Decimal,('154','308','416')))
        assert [a.audit.shipment.product_raw for a in saved.calculation.event_audit]==['Стандарт']*3
        assert [a.audit.eligible for a in saved.calculation.event_audit]==[True,False,False]
        assert saved.calculation.quarters[0].calculated_premium==Decimal('30')
    with WorkspaceRepository(path) as repo:
        service=WorkspaceService(repo)
        assert [kind for _,_,kind in service.product_mappings(employee.id,2026)]==['alias']*3
        assert repo.workspace_snapshot(saved.id).calculation==saved.calculation


def test_old_name_only_alias_does_not_cover_other_packaging():
    product=Product('p','Иннованта Стандарт 200 мл','ЭП','154','154')
    matcher=ProductMatcher([product],{'Стандарт':'p'})
    assert matcher.match('Стандарт')[0]==product
    assert matcher.match(product_mapping_key('Стандарт','500 мл'))==(None,None)
    assert matcher.match(product_mapping_key('Стандарт','200 мл'))==(None,None)


@pytest.mark.parametrize('source,target',[('200 мл','500 мл'),('200 мл','1000 мл'),('350 г','350 мл')])
def test_alias_with_conflicting_packaging_is_not_accepted(source,target):
    product=Product('p',f'Стандарт {target}','ЭП','1','1')
    key=product_mapping_key('Название из продаж',source)
    assert ProductMatcher([product],{key:'p'}).match(key)==(None,None)


def test_manual_mapping_rejects_wrong_volume(tmp_path):
    with WorkspaceRepository(tmp_path/'app.db') as repo:
        service=WorkspaceService(repo);service.login_admin('stopp')
        price=service.import_price(price_variants(tmp_path/'price.xlsx'))
        with pytest.raises(CalculationInputError,match='фасовк'):
            service.map_product('Стандарт',price.products[1].id,packaging='200 мл')
        assert not repo.aliases()


def test_litres_and_millilitres_share_alias_but_not_mass():
    assert normalize_text(product_mapping_key('Стандарт','1 л'))==normalize_text(product_mapping_key('Стандарт','1000 мл'))
    product=Product('p','Стандарт 1000 мл','ЭП','1','1')
    matcher=ProductMatcher([product])
    assert matcher.match(product_mapping_key(product.canonical_name,'1 л'))==(product,'exact')
    assert matcher.match(product_mapping_key(product.canonical_name,'1 кг'))==(None,None)


def test_fifo_does_not_allocate_return_to_different_packaging():
    product=Product('p','Стандарт 200 мл','ЭП','1','1')
    source=ShipmentEvent('sale','same-line','sheet',1,'Трофимов',2026,1,'Стандарт','10','1000',packaging='200 мл')
    refund=ReturnEvent('refund','same-line','sheet',1,'Трофимов',2026,7,'Стандарт','2','200',packaging='500 мл')
    calc=calculate_events([source],[refund],[product],Plans('0','0','0','0'),RolePolicy.kam(),
        aliases={product_mapping_key('Стандарт','200 мл'):'p'})
    assert calc.unallocated_returns==('refund',)
    assert not calc.allocations and calc.quarters[2].actual==-200


def test_old_event_snapshot_restores_with_empty_packaging():
    event=ShipmentEvent('s','l','sheet',1,'Трофимов',2026,1,'Товар','1','10')
    payload=json.loads(dump_object(event));payload['value']['fields'].pop('packaging')
    payload['value']['fields'].pop('legacy_line_key')
    assert load_object(json.dumps(payload))==event


def test_package_event_hashes_survive_reordering_and_returns_are_distinct(tmp_path):
    first=import_sales(packaged_sales(tmp_path/'one.xlsx',returns={7:('2','200')}),2026)
    moved=import_sales(packaged_sales(tmp_path/'two.xlsx',sizes=(1000,200,500),returns={7:('2','200')}),2026)
    for name in ('shipments','returns'):
        assert {e.packaging:e.event_id for e in getattr(first,name)}=={e.packaging:e.event_id for e in getattr(moved,name)}


def test_converted_legacy_event_ids_survive_package_row_reordering(tmp_path):
    path=packaged_sales(tmp_path/'legacy.xlsx')
    book=load_workbook(path);sheet=book[SHEETS[0]]
    column=next(c.column for c in sheet[6] if c.value=='Фасовка, мл')
    for row in range(7,10):sheet.cell(row,column).value=None
    book.save(path);book.close()
    with WorkspaceRepository(tmp_path/'app.db') as repo:
        service=WorkspaceService(repo)
        old=service.import_sales(path,2026)
        packaged=service.import_sales(packaged_sales(tmp_path/'first.xlsx'),2026)
        assert [e.event_id for e in packaged.shipments]==[e.event_id for e in old.shipments]
        reordered=service.import_sales(packaged_sales(tmp_path/'moved.xlsx',sizes=(1000,200,500)),2026)
        assert {e.packaging:e.event_id for e in reordered.shipments}=={e.packaging:e.event_id for e in packaged.shipments}


def test_package_reimport_keeps_legacy_paid_snapshot_and_deduplicates_return(tmp_path):
    with WorkspaceRepository(tmp_path/'app.db') as repo:
        service=WorkspaceService(repo);service.login_admin('stopp')
        employee=repo.resolve_employee('Трофимов');service.set_role(employee.id,2026,'KAM')
        service.save_inputs(employee.id,2026,Plans('0','0','0','0'))
        price=service.import_price(price_variants(tmp_path/'price.xlsx'))
        service.import_sales(sales_book(tmp_path/'old.xlsx',raw='Стандарт'),2026)
        service.map_product('Стандарт',price.products[0].id)
        old=service.calculate(employee.id,2026);service.confirm_payment(old.id,1)
        frozen=repo.workspace_snapshot(old.id)
        event_id=old.calculation.event_audit[0].event.event_id
        path=packaged_sales(tmp_path/'updated.xlsx',sizes=(200,),returns={7:('2','200')})
        first=service.import_sales(path,2026)
        assert first.shipments[0].event_id==event_id
        service.calculate(employee.id,2026)
        repeated=service.import_sales(path,2026)
        assert repeated==first
        service.calculate(employee.id,2026)
        assert repo.connection.execute('SELECT count(*) FROM return_events').fetchone()[0]==1
        assert repo.workspace_snapshot(old.id)==frozen
        assert repo.payments(employee.id,2026)[0].snapshot.calculation==old.calculation
        service.import_sales(packaged_sales(tmp_path/'changed.xlsx',sizes=(500,)),2026)
        service.map_product('Стандарт',price.products[1].id,packaging='500 мл')
        with pytest.raises(CalculationInputError,match='оплаченн'):
            service.calculate(employee.id,2026)


def test_numeric_packaging_retains_explicit_powder_unit_on_mixed_sheet(tmp_path):
    from app.models.import_data import SheetLayout
    from app.services.sales_importer import FIELDS
    path=packaged_sales(tmp_path/'mixed.xlsx',sizes=(350,))
    book=load_workbook(path);sheet=book[SHEETS[0]]
    product_col=next(c.column for c in sheet[6] if c.value=='Наименование')
    sheet.cell(7,product_col,'Иннованта Полуэлементная 350 г')
    columns={field:next(c.column for c in sheet[6] if c.value==label or field=='packaging' and c.value=='Фасовка, мл') for field,label in FIELDS.items()}
    book.save(path);book.close()
    imported=import_sales(path,2026,column_mappings={SHEETS[0]:SheetLayout(SHEETS[0],6,columns,'мл')})
    assert not imported.errors
    assert imported.shipments[0].packaging=='350 г'


@pytest.mark.parametrize('previous',['Innovanta_20102026','custom-password'])
def test_previous_default_password_upgrades_but_custom_password_survives(tmp_path,previous,caplog):
    path=tmp_path/'app.db'
    with WorkspaceRepository(path) as repo:
        salt=b'previous-unique-salt'
        digest=hashlib.pbkdf2_hmac('sha256',previous.encode(),salt,600_000)
        with repo.connection:
            repo.connection.execute('INSERT INTO admin_credentials VALUES(1,?,?,?)',(salt,digest,600_000))
        auth=AdminAuth(repo)
        assert auth.verify('stopp' if previous=='Innovanta_20102026' else previous)
        if previous=='Innovanta_20102026':assert not auth.verify(previous)
    assert b'stopp' not in path.read_bytes()
    assert previous.encode() not in path.read_bytes()
    assert 'stopp' not in caplog.text and previous not in caplog.text
    with WorkspaceRepository(path) as repo:
        assert AdminAuth(repo).verify('stopp' if previous=='Innovanta_20102026' else previous)


def test_password_can_be_changed_to_requested_short_password(tmp_path):
    with WorkspaceRepository(tmp_path/'app.db') as repo:
        auth=AdminAuth(repo);assert auth.login('stopp')
        auth.change_password('stopp','different-password')
        auth.change_password('different-password','stopp')
        assert auth.verify('stopp')
