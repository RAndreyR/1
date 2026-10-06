from decimal import Decimal
from hashlib import sha256
from pathlib import Path
import pytest
from app.services.excel_importer import ExcelImportError, import_excel, canonical_product_id
from tests.excel_helpers import make_workbook, patch_xml, SHIP_HEADERS, PRICE_COLUMNS
from tests.workbook_fixture import FIXTURE


@pytest.mark.parametrize('suffix', ['.xlsx', '.xlsm', '.XLSX'])
def test_supported_workbook_and_source_unchanged(tmp_path, suffix):
    path = make_workbook(tmp_path / ('input' + suffix))
    before = path.read_bytes()
    result = import_excel(path)
    assert not result.errors
    assert result.source_hash == sha256(before).hexdigest()
    assert result.source_file == str(path.resolve())
    assert path.read_bytes() == before
    assert result.shipment_row_count == len(result.shipments) == 1
    assert result.price_row_count == len(result.products) == 1
    assert result.shipments[0].source_row == 3
    assert result.shipments[0].revenue == Decimal('1000')
    assert result.shipments[0].client == 'Клиент'
    assert result.shipments[0].legal_entity == 'ООО'
    assert result.shipments[0].bitrix_task == 'task'
    assert result.shipments[0].comment == 'comment'
    assert result.products[0].category == 'СБКС'
    assert not hasattr(result, 'plans')


def test_headers_and_sheets_normalized_columns_reordered(tmp_path):
    headers = ['  ' + h.upper().replace(' ', '\u00a0  ') + ' ' for h in SHIP_HEADERS]
    prices = [' Наименование продукта ', ' Категория ', ' Зелёная зона для ЛПУ ', ' Зеленая зона для дистрибьюторов ']
    path = make_workbook(tmp_path/'input.xlsx', ship_headers=headers, price_headers=prices,
                         header_row=7, reorder=True, names=('  ОТГРУЗКИ\u00a0', ' Прайс '))
    result = import_excel(path)
    assert not result.errors
    assert result.shipments[0].revenue == 1000
    assert result.shipments[0].source_row == 8
    assert result.layouts[0].columns['revenue'] == 5
    assert result.layouts[0].header_row == 7
    assert result.products[0].price_lpu == 100


def test_ids_survive_price_list_reorder_and_case(tmp_path):
    a = import_excel(make_workbook(tmp_path/'a.xlsx', products=[['Товар', 'СБКС', 100, 100], ['Другой', 'ЭП', 90, 80]]))
    b = import_excel(make_workbook(tmp_path/'b.xlsx', products=[['Другой', 'ЭП', 90, 80], [' ТОВАР\u00a0', 'ВМК', 200, 150]]))
    assert a.products[0].id == b.products[1].id == canonical_product_id('Товар')
    assert b.products[1].category == 'ВМК'
    assert b.products[1].price_distributor == 150


@pytest.mark.parametrize('value,expected', [('1 234,56','1234.56'), ('1\u00a0234,56','1234.56'), ('-100,01','-100.01'), ('1.234e3','1234'), (0,'0')])
def test_numeric_text_and_signed_revenue(tmp_path, value, expected):
    rows = [['1 кв.', 'Клиент', 'ЛПУ', 'Товар', 10, value, None, None, None, None]]
    result = import_excel(make_workbook(tmp_path/'input.xlsx', shipments=rows))
    assert not result.errors
    assert result.shipments[0].revenue == Decimal(expected)
    assert result.shipments[0].unit_price is None


@pytest.mark.parametrize('value', [None, 'текст', 'NaN', 'Infinity', True, '1 2,34', '1,234.56'])
def test_bad_required_number_blocks_not_silently_zero(tmp_path, value):
    rows = [['1 кв.', 'Клиент', 'ЛПУ', 'Товар', 10, value, None, None, None, None]]
    result = import_excel(make_workbook(tmp_path/'input.xlsx', shipments=rows))
    assert len(result.shipments) == 0
    assert result.shipment_row_count == 1
    issue = next(i for i in result.errors if i.code == 'invalid_number')
    assert issue.source_row == 3 and issue.field == 'revenue'


def test_raw_numeric_xml_precision_and_cached_formula(tmp_path):
    path = make_workbook(tmp_path/'input.xlsx')
    patch_xml(path, 1, 'F3', token='1234567.890123456789012345')
    patch_xml(path, 1, 'G3', token='154.049999999999999999', formula='F3/E3')
    result = import_excel(path)
    assert not result.errors
    assert result.shipments[0].revenue == Decimal('1234567.890123456789012345')
    assert result.shipments[0].unit_price == Decimal('154.049999999999999999')


@pytest.mark.parametrize('coordinate,blocking', [('F3',True), ('E3',True), ('G3',False)])
def test_uncached_formula_reported_without_execution(tmp_path, coordinate, blocking):
    path = make_workbook(tmp_path/'input.xlsx')
    patch_xml(path, 1, coordinate, formula='SUM(1000,1000)')
    result = import_excel(path)
    issue = next(i for i in result.issues if i.code == 'uncached_formula')
    assert issue.severity == ('error' if blocking else 'warning')
    assert bool(result.errors) == blocking
    if not blocking:
        assert result.shipments[0].unit_price is None


def test_excel_error_is_not_a_price(tmp_path):
    path = make_workbook(tmp_path/'input.xlsx')
    patch_xml(path, 1, 'G3', token='#DIV/0!', kind='e', formula='F3/E3')
    result = import_excel(path)
    assert result.errors[0].code == 'invalid_number'
    assert result.errors[0].field == 'unit_price'


def test_missing_and_duplicate_columns(tmp_path):
    headers = list(SHIP_HEADERS)
    headers[-1] = 'Количество'
    result = import_excel(make_workbook(tmp_path/'input.xlsx', ship_headers=headers))
    assert {i.code for i in result.errors} == {'missing_column','duplicate_column'}
    assert not result.shipments
    assert result.shipment_row_count == 1


def test_missing_sheet_and_unknown_headers(tmp_path):
    path = make_workbook(tmp_path/'input.xlsx', names=('Данные','прайс'), price_headers=['a','b','c','d'])
    result = import_excel(path)
    assert {i.code for i in result.errors} == {'missing_sheet','header_not_found'}
    assert result.detected_sheets == ('Данные', 'прайс', 'KPI')


def test_ambiguous_normalized_sheets(tmp_path):
    from openpyxl import load_workbook
    path = make_workbook(tmp_path/'input.xlsx')
    wb = load_workbook(path)
    wb.create_sheet(' отгрузки ')
    wb.save(path)
    wb.close()
    result = import_excel(path)
    assert any(i.code == 'ambiguous_sheet' for i in result.errors)


def test_empty_sheets_and_blank_rows(tmp_path):
    result = import_excel(make_workbook(tmp_path/'input.xlsx', shipments=[[None]*10], products=[]))
    assert len(result.errors) == 2
    assert all(i.code == 'empty_sheet' for i in result.errors)
    assert result.shipment_row_count == 0


@pytest.mark.parametrize('products,code', [
    ([['Товар','СБКС',100,100], [' ТОВАР ','ЭП',100,100]], 'duplicate_product'),
    ([['Товар','Непонятно',100,100]], 'invalid_product'),
    ([[None,'СБКС',100,100]], 'invalid_product'),
    ([['Товар','СБКС',-1,100]], 'invalid_product'),
    ([['Товар','СБКС','abc',100]], 'invalid_number'),
])
def test_invalid_price_list(tmp_path, products, code):
    result = import_excel(make_workbook(tmp_path/'input.xlsx', products=products))
    assert any(i.code == code for i in result.errors)


def test_blank_zero_thresholds_allowed_for_audit(tmp_path):
    result = import_excel(make_workbook(tmp_path/'input.xlsx', products=[['Товар','СБКС',None,0]]))
    assert not result.errors
    assert result.products[0].price_lpu is None
    assert result.products[0].price_distributor == 0


def test_stale_dimensions_do_not_truncate_import(tmp_path):
    path = make_workbook(tmp_path/'input.xlsx')
    patch_xml(path, 1, 'F3', token='1000', dimension='A1:A1')
    result = import_excel(path)
    assert not result.errors and len(result.shipments) == 1


@pytest.mark.parametrize('suffix', ['.xls', '.csv'])
def test_unsupported_format(tmp_path, suffix):
    with pytest.raises(ExcelImportError, match='.xlsx'):
        import_excel(tmp_path/('input'+suffix))


def test_corrupt_or_missing_file(tmp_path):
    path = tmp_path/'input.xlsx'
    with pytest.raises(ExcelImportError):
        import_excel(path)
    path.write_bytes(b'not an Excel archive')
    with pytest.raises(ExcelImportError):
        import_excel(path)


def test_real_macro_workbook_is_read_only():
    before = FIXTURE.read_bytes()
    result = import_excel(FIXTURE)
    assert not result.errors
    assert len(result.shipments) == 170 and len(result.products) == 28
    assert FIXTURE.read_bytes() == before


def test_normal_openpyxl_formula_with_empty_cache(tmp_path):
    rows = [['1 кв.', 'Клиент', 'ЛПУ', 'Товар', 10, 1000, '=F3/E3', None, None, None]]
    result = import_excel(make_workbook(tmp_path/'input.xlsx', shipments=rows))
    assert not result.errors
    assert result.shipments[0].unit_price is None
    assert any(i.code == 'uncached_formula' and i.field == 'unit_price' for i in result.issues)


def test_date_cannot_be_treated_as_revenue(tmp_path):
    from datetime import datetime
    rows = [['1 кв.', 'Клиент', 'ЛПУ', 'Товар', 10, datetime(2026,1,1), None, None, None, None]]
    result = import_excel(make_workbook(tmp_path/'input.xlsx', shipments=rows))
    assert any(i.code == 'invalid_number' and i.field == 'revenue' for i in result.errors)
