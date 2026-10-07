from datetime import datetime
from openpyxl import Workbook
from app.services.sales_importer import FIELDS,SHEETS


def sales_book(path,*,manager='Трофимов',months=None,returns=None,year=2026,shift=0,
               missing_headers=False,db='Дистрибьютер',raw='Товар',contract='K-1'):
    workbook=Workbook();workbook.remove(workbook.active)
    months=months if months is not None else {1:('10','1000')}
    returns=returns or {}
    for index,name in enumerate(SHEETS):
        sheet=workbook.create_sheet(name)
        group_row=5+shift;header=group_row+1
        fields=list(FIELDS)
        # Reverse metadata order and offset monthly groups: no fixed Excel letters.
        fields.reverse()
        values={'fo':'ФО','region':'Регион','manager':manager,'lpu':'Больница','db':db,
                'legal_entity':'ЮЛ','contract':contract,'product_raw':raw}
        for col,field in enumerate(fields,2):
            if not missing_headers or index==0:
                sheet.cell(header,col,FIELDS[field])
            sheet.cell(header+1,col,values[field])
        col=14+shift
        for month,pair in months.items():
            sheet.cell(group_row,col,datetime(year,month,1));sheet.merge_cells(start_row=group_row,start_column=col,end_row=group_row,end_column=col+1)
            # Deliberately reversed amount/quantity for signed header-based lookup.
            sheet.cell(header,col,'сумма');sheet.cell(header,col+1,'кг')
            sheet.cell(header+1,col,pair[1] if index==0 else '0');sheet.cell(header+1,col+1,pair[0] if index==0 else '0')
            col+=3
        month_names=['январь','февраль','март','апрель','май','июнь','июль','август','сентябрь','октябрь','ноябрь','декабрь']
        for month,pair in returns.items():
            sheet.cell(group_row,col,'Возврат '+month_names[month-1]);sheet.cell(group_row,col+1,'Возврат '+month_names[month-1])
            sheet.cell(header+1,col,pair[0] if index==0 else '0');sheet.cell(header+1,col+1,pair[1] if index==0 else '0');col+=3
    embedded=workbook.create_sheet('прайс');embedded.append(['Нельзя использовать встроенный прайс'])
    workbook.save(path);return path


def price_book(path,*,category='ВМК',threshold='100'):
    workbook=Workbook();sheet=workbook.active
    sheet.append(['Продукт','Группа','Зеленая зона ЛПУ','Зеленая зона дистрибьюторов'])
    sheet.append(['Товар',category,threshold,threshold]);workbook.save(path);return path
